#!/usr/bin/env python3
# SPDX-License-Identifier: CC0-1.0

# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "numpy==2.5.1",
#   "openpyxl==3.1.5",
#   "pybigwig==0.3.25",
# ]
# ///

"""Prepare and validate the mouse fetal-development GenomeSpy dataset."""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
import math
import re
import shutil
import time
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import numpy  # type: ignore[import-not-found]
import pyBigWig  # type: ignore[import-not-found]
from openpyxl import load_workbook  # type: ignore[import-untyped]

RECIPE_DIR = Path(__file__).resolve().parents[1]
PROVENANCE_PATH = RECIPE_DIR / "provenance.json"
DOWNLOAD_DIR = RECIPE_DIR / "download"
OUTPUT_DIR = RECIPE_DIR / "output"
WORK_DIR = RECIPE_DIR / "work"
RNA_DIR = DOWNLOAD_DIR / "rna"
SUPPLEMENT_DIR = DOWNLOAD_DIR / "supplements"
BIGWIG_DIR = OUTPUT_DIR / "bigwigs"
ACCEPTED_RUN_PATH = WORK_DIR / "accepted-run.json"
USER_AGENT = "GenomeSpy dataset recipe encode-mouse-fetal-development-mm10"

CHROM_LENGTHS = {
    "chr10": 130694993,
    "chr12": 120129022,
    "chr7": 145441459,
    "chr9": 124595110,
}
TISSUE_ORDER = {"Forebrain": 0, "Heart": 1, "Limb": 2}


@dataclass(frozen=True)
class FileLock:
    """One pinned ENCODE input file."""

    accession: str
    experiment: str
    tissue: str
    stage_days: float
    stage_label: str
    biological_replicate: int
    size: int
    md5: str
    url: str
    technical_replicate: str = ""
    biosample: str = ""

    @property
    def sample_id(self) -> str:
        tissue = self.tissue.lower()
        stage = str(self.stage_days).replace(".", "-")
        return f"{tissue}-e{stage}-rep{self.biological_replicate}"

    @property
    def condition(self) -> tuple[str, float]:
        return (self.tissue, self.stage_days)


@dataclass(frozen=True)
class Region:
    """A retained zero-based, half-open interval."""

    region_id: str
    chrom: str
    start: int
    end: int
    story: str


@dataclass(frozen=True)
class Gene:
    """A selected expression gene."""

    symbol: str
    gene_id: str


@dataclass(frozen=True)
class ExpressionValue:
    """One selected gene's TPM in one RNA replicate."""

    lock: FileLock
    gene: Gene
    tpm: float


ELEMENTS: tuple[tuple[str, str, int, int, str, str, str, str], ...] = (
    (
        "Ascl1 prediction 1",
        "chr10",
        87308700,
        87310700,
        "Forebrain",
        "enhancer-gene prediction",
        "Ascl1",
        "Supplementary Table 8c; supported in both prediction replicates",
    ),
    (
        "Ascl1 prediction 2",
        "chr10",
        87350700,
        87352700,
        "Forebrain",
        "enhancer-gene prediction",
        "Ascl1",
        "Supplementary Table 8c; supported in both prediction replicates",
    ),
    (
        "Ascl1 prediction 3",
        "chr10",
        87446400,
        87448800,
        "Forebrain",
        "enhancer-gene prediction",
        "Ascl1",
        "Supplementary Table 8c; supported in both prediction replicates",
    ),
    (
        "Ascl1 prediction 4",
        "chr10",
        87472500,
        87474500,
        "Forebrain",
        "enhancer-gene prediction",
        "Ascl1",
        "Supplementary Table 8c; supported in both prediction replicates",
    ),
    (
        "Ascl1 prediction 5",
        "chr10",
        87487400,
        87491200,
        "Forebrain",
        "enhancer-gene prediction",
        "Ascl1",
        "Supplementary Table 8c; supported in both prediction replicates",
    ),
    (
        "mEN886 / mm1606",
        "chr12",
        111691971,
        111695499,
        "Forebrain",
        "transgenic reporter",
        "",
        (
            "Supplementary Table 10: forebrain positive 6/6; midbrain, "
            "hindbrain, and neural tube 6/6"
        ),
    ),
    (
        "Ckb prediction overlapping mEN886",
        "chr12",
        111690600,
        111696400,
        "Forebrain",
        "enhancer-gene prediction",
        "Ckb",
        "Supplementary Table 8c; supported in both prediction replicates",
    ),
    (
        "mEN978 / mm1683",
        "chr7",
        139466376,
        139469610,
        "Heart",
        "transgenic reporter",
        "",
        "Supplementary Table 10: heart positive 5/8",
    ),
    (
        "mEN918 / mm1617",
        "chr9",
        43252109,
        43255189,
        "Limb",
        "transgenic reporter",
        "",
        "Supplementary Table 10: limb positive 4/4; somite 3/4",
    ),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--verify-only", action="store_true")
    return parser.parse_args()


def load_provenance() -> dict[str, Any]:
    value: Any = json.loads(PROVENANCE_PATH.read_text(encoding="utf-8"))
    if not isinstance(value, dict) or value.get("recipeId") != RECIPE_DIR.name:
        raise ValueError("provenance.json does not describe this recipe")
    return value


def source(provenance: dict[str, Any], source_id: str) -> dict[str, Any]:
    matches = [item for item in provenance["sources"] if item["id"] == source_id]
    if len(matches) != 1:
        raise ValueError(f"Expected one source {source_id!r}")
    return matches[0]


def file_locks(value: dict[str, Any]) -> list[FileLock]:
    locks = [
        FileLock(
            accession=str(item["accession"]),
            experiment=str(item["experiment"]),
            tissue=str(item["tissue"]),
            stage_days=float(item["stageDays"]),
            stage_label=str(item["stageLabel"]),
            biological_replicate=int(item["biologicalReplicate"]),
            size=int(item["fileSizeBytes"]),
            md5=str(item["md5"]),
            url=str(item["cloudUrl"]),
            technical_replicate=str(item.get("technicalReplicate", "")),
            biosample=str(item.get("biosample", "")),
        )
        for item in value["files"]
    ]
    expected = sorted(
        locks,
        key=lambda item: (
            TISSUE_ORDER[item.tissue],
            item.stage_days,
            item.biological_replicate,
        ),
    )
    if locks != expected:
        raise ValueError(
            "Accepted inputs must use tissue, numeric stage, replicate order"
        )
    keys = [(item.tissue, item.stage_days, item.biological_replicate) for item in locks]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate biological-replicate condition key")
    return locks


def regions(provenance: dict[str, Any]) -> list[Region]:
    values = [
        Region(
            str(item["id"]),
            str(item["chrom"]),
            int(item["start"]),
            int(item["end"]),
            str(item["story"]),
        )
        for item in provenance["parameters"]["regions"]
    ]
    if len(values) != 4 or any(item.end <= item.start for item in values):
        raise ValueError("Expected four nonempty retained regions")
    if any(item.end - item.start != 1_000_000 for item in values):
        raise ValueError("Every retained region must provide one megabase of context")
    return values


def genes(provenance: dict[str, Any]) -> list[Gene]:
    values = [
        Gene(str(item["symbol"]), str(item["geneId"]))
        for item in provenance["parameters"]["genes"]
    ]
    if len(values) != 8 or len({item.symbol for item in values}) != 8:
        raise ValueError("Expected eight unique expression genes")
    return values


def digest(path: Path, algorithm: str = "sha256") -> str:
    value = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def validate_file(path: Path, size: int, md5: str, sha256: str | None = None) -> None:
    if not path.is_file() or path.stat().st_size != size:
        raise ValueError(f"Pinned size mismatch: {path}")
    if digest(path, "md5") != md5:
        raise ValueError(f"Pinned MD5 mismatch: {path}")
    if sha256 is not None and digest(path) != sha256:
        raise ValueError(f"Pinned SHA-256 mismatch: {path}")


def download(url: str, destination: Path, size: int, md5: str) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".part")
    for attempt in range(5):
        offset = temporary.stat().st_size if temporary.exists() else 0
        headers = {"User-Agent": USER_AGENT}
        if offset:
            headers["Range"] = f"bytes={offset}-"
        try:
            with urlopen(Request(url, headers=headers), timeout=120) as response:
                resume = offset > 0 and response.status == 206
                with temporary.open("ab" if resume else "wb") as output:
                    shutil.copyfileobj(response, output, 4 * 1024 * 1024)
            break
        except (HTTPError, URLError, TimeoutError):
            if attempt == 4:
                raise
            time.sleep(2**attempt)
    validate_file(temporary, size, md5)
    temporary.replace(destination)


def ensure_download(value: dict[str, Any], destination: Path) -> Path:
    if not destination.exists():
        download(
            str(value["url"] if "url" in value else value["cloudUrl"]),
            destination,
            int(value["fileSizeBytes"]),
            str(value["md5"]),
        )
    validate_file(
        destination,
        int(value["fileSizeBytes"]),
        str(value["md5"]),
        str(value["sha256"]) if "sha256" in value else None,
    )
    return destination


def prepare_inputs(
    provenance: dict[str, Any], rna: list[FileLock]
) -> tuple[Path, Path]:
    gencode = source(provenance, "gencode-m21")
    gtf_path = ensure_download(
        gencode,
        DOWNLOAD_DIR
        / "annotation"
        / "gencode.vM21.primary_assembly.annotation_UCSC_names.gtf.gz",
    )
    for lock in rna:
        ensure_download(
            {
                "cloudUrl": lock.url,
                "fileSizeBytes": lock.size,
                "md5": lock.md5,
            },
            RNA_DIR / f"{lock.accession}.tsv",
        )
    supplement = source(provenance, "gorkin-supplementary-tables")
    supplement_paths: dict[str, Path] = {}
    for item in supplement["files"]:
        supplement_paths[str(item["localName"])] = ensure_download(
            item, SUPPLEMENT_DIR / str(item["localName"])
        )
    return gtf_path, supplement_paths["MOESM6.xlsx"]


def validate_supplement(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    required_sheets = {
        "S8c.Enhancer-Gene-Map-Replicatd",
        "S10.enhancer_validation_results",
    }
    if not required_sheets.issubset(workbook.sheetnames):
        raise ValueError("Gorkin workbook no longer has the required worksheets")
    s8 = {
        (str(row[0]), int(row[1]), int(row[2]), str(row[4]), str(row[8]))
        for row in workbook["S8c.Enhancer-Gene-Map-Replicatd"].iter_rows(
            values_only=True
        )
        if row[0] in {"chr10", "chr12"} and isinstance(row[1], int)
    }
    for element in ELEMENTS[:5]:
        if (element[1], element[2], element[3], "Ascl1", "Ascl1") not in s8:
            raise ValueError(
                f"Ascl1 prediction changed in source workbook: {element[0]}"
            )
    if ("chr12", 111690600, 111696400, "Ckb", "Ckb") not in s8:
        raise ValueError("Ckb prediction overlapping mEN886 changed")
    s10_rows = {
        (str(row[2]), str(row[3]), str(row[4]), str(row[5]), str(row[6] or ""))
        for row in workbook["S10.enhancer_validation_results"].iter_rows(
            values_only=True
        )
        if row[2] is not None
    }
    expected = {
        (
            "mm1606",
            "mEN886",
            "chr12:111691971-111695499",
            "Fb positive (6/6)",
            "Mb (6/6), Hb (6/6), neural tube (6/6)",
        ),
        ("mm1683", "mEN978", "chr7:139466376-139469610", "Ht positive (5/8)", ""),
        (
            "mm1617",
            "mEN918",
            "chr9:43252109-43255189",
            "Lb positive (4/4)",
            "somite (3/4)",
        ),
    }
    if not expected.issubset(s10_rows):
        raise ValueError("Selected reporter records changed in Supplementary Table 10")


def parse_attributes(value: str) -> dict[str, str]:
    return dict(re.findall(r'(\w+) "([^"]+)"', value))


def read_gene_annotation(
    gtf_path: Path, retained_regions: list[Region], selected_genes: list[Gene]
) -> list[tuple[str, int, int, str, str, str]]:
    expected = {(item.symbol, item.gene_id) for item in selected_genes}
    observed: set[tuple[str, str]] = set()
    rows: list[tuple[str, int, int, str, str, str]] = []
    by_chrom = defaultdict(list)
    for region in retained_regions:
        by_chrom[region.chrom].append(region)
    with gzip.open(gtf_path, "rt", encoding="utf-8") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            if len(fields) != 9 or fields[2] != "gene":
                continue
            attrs = parse_attributes(fields[8])
            pair = (attrs.get("gene_name", ""), attrs.get("gene_id", ""))
            if pair in expected:
                observed.add(pair)
            chrom = fields[0]
            start = int(fields[3]) - 1
            end = int(fields[4])
            if any(
                start < region.end and end > region.start for region in by_chrom[chrom]
            ):
                rows.append((chrom, start, end, fields[6], pair[1], pair[0]))
    if observed != expected:
        raise ValueError(f"Selected GENCODE mappings changed: {observed!r}")
    return sorted(rows, key=lambda row: (row[0], row[1], row[2], row[4]))


def read_expression(
    locks: list[FileLock], selected_genes: list[Gene]
) -> list[ExpressionValue]:
    by_id = {item.gene_id: item for item in selected_genes}
    output: list[ExpressionValue] = []
    for lock in locks:
        found: dict[str, float] = {}
        with (RNA_DIR / f"{lock.accession}.tsv").open(
            encoding="utf-8", newline=""
        ) as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            if reader.fieldnames is None or not {"gene_id", "TPM"}.issubset(
                reader.fieldnames
            ):
                raise ValueError(f"RNA table lacks gene_id or TPM: {lock.accession}")
            for row in reader:
                gene_id = row["gene_id"]
                if gene_id not in by_id:
                    continue
                tpm = float(row["TPM"])
                if not math.isfinite(tpm) or tpm < 0 or gene_id in found:
                    raise ValueError(f"Invalid or duplicate TPM in {lock.accession}")
                found[gene_id] = tpm
        if set(found) != set(by_id):
            raise ValueError(f"Selected genes missing from {lock.accession}")
        output.extend(
            ExpressionValue(lock, gene, found[gene.gene_id]) for gene in selected_genes
        )
    return output


def expression_summaries(
    values: list[ExpressionValue],
) -> dict[tuple[str, float, str], dict[str, Any]]:
    grouped: dict[tuple[str, float, str], list[ExpressionValue]] = defaultdict(list)
    for value in values:
        grouped[(value.lock.tissue, value.lock.stage_days, value.gene.symbol)].append(
            value
        )
    output: dict[tuple[str, float, str], dict[str, Any]] = {}
    for key, group in grouped.items():
        if len(group) != 2:
            raise ValueError(f"Expected two RNA replicates for {key!r}")
        mean = float(numpy.mean([item.tpm for item in group]))
        output[key] = {
            "meanTpm": mean,
            "log2MeanTpmPlus1": math.log2(mean + 1),
            "experiment": group[0].lock.experiment,
            "files": [item.lock.accession for item in group],
            "replicateCount": len(group),
        }
    for gene in {key[2] for key in output}:
        keys = [key for key in output if key[2] == gene]
        transformed = numpy.array(
            [output[key]["log2MeanTpmPlus1"] for key in keys], dtype=float
        )
        if len(keys) != 12:
            raise ValueError(f"Expected 12 tissue-stage conditions for {gene}")
        standard_deviation = float(numpy.std(transformed))
        if not math.isfinite(standard_deviation) or standard_deviation <= 0:
            raise ValueError(f"Cannot standardize expression for {gene}")
        mean = float(numpy.mean(transformed))
        for key, value in zip(keys, transformed, strict=True):
            output[key]["zScore"] = (float(value) - mean) / standard_deviation
    return output


def extract_bigwig(lock: FileLock, retained_regions: list[Region]) -> None:
    destination = BIGWIG_DIR / f"{lock.sample_id}.bigWig"
    if destination.exists():
        return
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".part")
    temporary.unlink(missing_ok=True)
    remote = pyBigWig.open(lock.url)
    if remote is None or not remote.isBigWig():
        raise ValueError(f"Could not open {lock.accession}")
    observed = remote.chroms()
    if {chrom: observed.get(chrom) for chrom in CHROM_LENGTHS} != CHROM_LENGTHS:
        remote.close()
        raise ValueError(f"Unexpected mm10 chromosome lengths: {lock.accession}")
    output = pyBigWig.open(str(temporary), "w")
    try:
        output.addHeader(list(CHROM_LENGTHS.items()))
        for region in retained_regions:
            intervals = remote.intervals(region.chrom, region.start, region.end) or ()
            clipped = [
                (max(int(start), region.start), min(int(end), region.end), float(score))
                for start, end, score in intervals
                if min(int(end), region.end) > max(int(start), region.start)
            ]
            if any(not math.isfinite(score) or score < 0 for _, _, score in clipped):
                raise ValueError(f"Invalid source signal in {lock.accession}")
            if clipped:
                output.addEntries(
                    [region.chrom] * len(clipped),
                    [item[0] for item in clipped],
                    ends=[item[1] for item in clipped],
                    values=[item[2] for item in clipped],
                )
    finally:
        output.close()
        remote.close()
    temporary.replace(destination)


def tsv_text(fields: Iterable[str], rows: Iterable[Iterable[Any]]) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter="\t", lineterminator="\n")
    writer.writerow(fields)
    writer.writerows(rows)
    return output.getvalue()


def format_number(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else format(value, ".10g")


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".part")
    temporary.write_text(value, encoding="utf-8")
    temporary.replace(path)


def write_tables(
    provenance: dict[str, Any],
    h3: list[FileLock],
    rna_values: list[ExpressionValue],
    selected_genes: list[Gene],
    retained_regions: list[Region],
    annotation_rows: list[tuple[str, int, int, str, str, str]],
) -> None:
    summaries = expression_summaries(rna_values)
    zscore_fields = [f"RNA-seq.{item.symbol}" for item in selected_genes]
    expression_fields = [
        f"RNA-seq.log2Abundance.{item.symbol}" for item in selected_genes
    ]
    mean_fields = [f"RNA-seq.meanTpm.{item.symbol}" for item in selected_genes]
    sample_fields = [
        "sampleId",
        "rowLabel",
        "Sample.tissue",
        "Sample.stageDays",
        "Sample.stageLabel",
        "H3K27ac.biologicalReplicate",
        "H3K27ac.technicalReplicate",
        "H3K27ac.biosample",
        "H3K27ac.experiment",
        "H3K27ac.fileAccession",
        "Sample.strain",
        "Sample.preparation",
        "H3K27ac.signal",
        "RNA-seq.provenance.join",
        "RNA-seq.provenance.experiment",
        "RNA-seq.provenance.replicateCount",
        "RNA-seq.provenance.fileAccessions",
        *zscore_fields,
        *expression_fields,
        *mean_fields,
    ]
    sample_rows: list[list[Any]] = []
    for lock in h3:
        first = summaries[(lock.tissue, lock.stage_days, selected_genes[0].symbol)]
        row: list[Any] = [
            lock.sample_id,
            f"{lock.tissue} {lock.stage_label} R{lock.biological_replicate}",
            lock.tissue,
            format_number(lock.stage_days),
            lock.stage_label,
            lock.biological_replicate,
            lock.technical_replicate,
            lock.biosample,
            lock.experiment,
            lock.accession,
            "B6NCrl",
            "dissected embryonic tissue pool",
            "H3K27ac fold change over control",
            "RNA-seq tissue-stage mean; shared across ChIP replicates",
            first["experiment"],
            first["replicateCount"],
            ",".join(first["files"]),
        ]
        row.extend(
            format_number(
                summaries[(lock.tissue, lock.stage_days, gene.symbol)]["zScore"]
            )
            for gene in selected_genes
        )
        row.extend(
            format_number(
                summaries[(lock.tissue, lock.stage_days, gene.symbol)][
                    "log2MeanTpmPlus1"
                ]
            )
            for gene in selected_genes
        )
        row.extend(
            format_number(
                summaries[(lock.tissue, lock.stage_days, gene.symbol)]["meanTpm"]
            )
            for gene in selected_genes
        )
        sample_rows.append(row)
    write_text(OUTPUT_DIR / "samples.tsv", tsv_text(sample_fields, sample_rows))

    replicate_rows = [
        (
            value.lock.tissue,
            format_number(value.lock.stage_days),
            value.lock.stage_label,
            value.lock.experiment,
            value.lock.accession,
            value.lock.biological_replicate,
            value.gene.gene_id,
            value.gene.symbol,
            format_number(value.tpm),
        )
        for value in rna_values
    ]
    write_text(
        OUTPUT_DIR / "expression-replicates.tsv",
        tsv_text(
            (
                "tissue",
                "stageDays",
                "stageLabel",
                "rnaExperiment",
                "rnaFileAccession",
                "rnaBiologicalReplicate",
                "geneId",
                "symbol",
                "tpm",
            ),
            replicate_rows,
        ),
    )
    condition_rows = []
    for tissue in TISSUE_ORDER:
        for stage in (11.5, 12.5, 13.5, 15.5):
            for gene in selected_genes:
                item = summaries[(tissue, stage, gene.symbol)]
                condition_rows.append(
                    (
                        tissue,
                        format_number(stage),
                        f"E{stage}",
                        item["experiment"],
                        gene.gene_id,
                        gene.symbol,
                        format_number(item["meanTpm"]),
                        format_number(item["log2MeanTpmPlus1"]),
                        format_number(item["zScore"]),
                        item["replicateCount"],
                        ",".join(item["files"]),
                    )
                )
    write_text(
        OUTPUT_DIR / "expression-conditions.tsv",
        tsv_text(
            (
                "tissue",
                "stageDays",
                "stageLabel",
                "rnaExperiment",
                "geneId",
                "symbol",
                "meanTpm",
                "log2MeanTpmPlus1",
                "zScore",
                "rnaReplicateCount",
                "rnaFileAccessions",
            ),
            condition_rows,
        ),
    )
    write_text(
        OUTPUT_DIR / "regions.tsv",
        tsv_text(
            ("chrom", "start", "end", "name", "description"),
            (
                (item.chrom, item.start, item.end, item.region_id, item.story)
                for item in retained_regions
            ),
        ),
    )
    write_text(
        OUTPUT_DIR / "elements.tsv",
        tsv_text(
            (
                "name",
                "chrom",
                "start",
                "end",
                "intendedTissue",
                "evidenceType",
                "predictedTarget",
                "sourceEvidence",
            ),
            ELEMENTS,
        ),
    )
    write_text(
        OUTPUT_DIR / "genes.tsv",
        tsv_text(
            ("chrom", "start", "end", "strand", "geneId", "symbol"), annotation_rows
        ),
    )
    assessments = provenance["parameters"]["candidateAssessment"]
    write_text(
        OUTPUT_DIR / "candidate-assessment.tsv",
        tsv_text(
            ("candidate", "decision", "e12_5SignalSummary", "reason"),
            (
                (
                    item["candidate"],
                    item["decision"],
                    item["e12_5SignalSummary"],
                    item["reason"],
                )
                for item in assessments
            ),
        ),
    )
    selection_rows = []
    for assay, locks in (
        ("H3K27ac", h3),
        ("RNA-seq", [value.lock for value in rna_values[:: len(selected_genes)]]),
    ):
        for lock in locks:
            selection_rows.append(
                (
                    assay,
                    "accepted",
                    lock.tissue,
                    format_number(lock.stage_days),
                    lock.experiment,
                    lock.accession,
                    lock.biological_replicate,
                    lock.size,
                    "one processed file per biological replicate",
                )
            )
    write_text(
        OUTPUT_DIR / "selection-report.tsv",
        tsv_text(
            (
                "assay",
                "decision",
                "tissue",
                "stageDays",
                "experiment",
                "fileAccession",
                "biologicalReplicate",
                "sourceFileSizeBytes",
                "reason",
            ),
            selection_rows,
        ),
    )


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def validate_outputs(
    h3: list[FileLock],
    rna: list[FileLock],
    selected_genes: list[Gene],
    retained_regions: list[Region],
) -> dict[str, Any]:
    samples = read_tsv(OUTPUT_DIR / "samples.tsv")
    conditions = read_tsv(OUTPUT_DIR / "expression-conditions.tsv")
    replicates = read_tsv(OUTPUT_DIR / "expression-replicates.tsv")
    if len(samples) != 24 or len({row["sampleId"] for row in samples}) != 24:
        raise ValueError("samples.tsv must contain 24 unique ChIP rows")
    if len({row["rowLabel"] for row in samples}) != 24:
        raise ValueError("samples.tsv must contain 24 unique row labels")
    expected_order = sorted(
        samples,
        key=lambda row: (
            TISSUE_ORDER[row["Sample.tissue"]],
            float(row["Sample.stageDays"]),
            int(row["H3K27ac.biologicalReplicate"]),
        ),
    )
    if samples != expected_order:
        raise ValueError("Sample ordering is not tissue, numeric stage, replicate")
    if len(conditions) != 96 or len(replicates) != 192:
        raise ValueError("Unexpected expression table dimensions")
    replicate_lookup: dict[tuple[str, str, str], list[float]] = defaultdict(list)
    for row in replicates:
        tpm = float(row["tpm"])
        if not math.isfinite(tpm) or tpm < 0:
            raise ValueError("Expression TPM must be finite and nonnegative")
        replicate_lookup[(row["tissue"], row["stageDays"], row["symbol"])].append(tpm)
    conditions_by_gene: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in conditions:
        values = replicate_lookup[(row["tissue"], row["stageDays"], row["symbol"])]
        if len(values) != 2 or not math.isclose(
            float(row["meanTpm"]), float(numpy.mean(values)), rel_tol=1e-9, abs_tol=1e-9
        ):
            raise ValueError("Condition mean does not trace to two RNA replicates")
        expected_log = math.log2(float(row["meanTpm"]) + 1)
        if not math.isclose(
            float(row["log2MeanTpmPlus1"]), expected_log, rel_tol=1e-9, abs_tol=1e-9
        ):
            raise ValueError("Expression transform is not log2(mean TPM + 1)")
        conditions_by_gene[row["symbol"]].append(row)
    for symbol, rows in conditions_by_gene.items():
        transformed = numpy.array(
            [float(row["log2MeanTpmPlus1"]) for row in rows], dtype=float
        )
        expected = (transformed - numpy.mean(transformed)) / numpy.std(transformed)
        observed = numpy.array([float(row["zScore"]) for row in rows], dtype=float)
        if len(rows) != 12 or not numpy.allclose(
            observed, expected, rtol=1e-8, atol=1e-8
        ):
            raise ValueError(f"Expression z-scores are invalid for {symbol}")
    condition_lookup = {
        (row["tissue"], row["stageDays"], row["symbol"]): row for row in conditions
    }
    for sample in samples:
        for gene in selected_genes:
            item = condition_lookup[
                (sample["Sample.tissue"], sample["Sample.stageDays"], gene.symbol)
            ]
            if (
                sample[f"RNA-seq.{gene.symbol}"] != item["zScore"]
                or sample[f"RNA-seq.log2Abundance.{gene.symbol}"]
                != item["log2MeanTpmPlus1"]
                or sample[f"RNA-seq.meanTpm.{gene.symbol}"] != item["meanTpm"]
            ):
                raise ValueError("Sample metadata does not match condition expression")

    total_size = 0
    maxima: list[float] = []
    for lock in h3:
        path = BIGWIG_DIR / f"{lock.sample_id}.bigWig"
        total_size += path.stat().st_size
        with pyBigWig.open(str(path)) as bigwig:
            if bigwig.chroms() != CHROM_LENGTHS:
                raise ValueError(f"Unexpected regional BigWig header: {path}")
            for region in retained_regions:
                if bigwig.intervals(region.chrom, 0, region.start) or bigwig.intervals(
                    region.chrom, region.end, CHROM_LENGTHS[region.chrom]
                ):
                    raise ValueError(f"Signal outside retained bounds: {path}")
                intervals = (
                    bigwig.intervals(region.chrom, region.start, region.end) or ()
                )
                if not intervals:
                    raise ValueError(f"No records in {region.region_id}: {path}")
                scores = [float(item[2]) for item in intervals]
                if any(not math.isfinite(item) or item < 0 for item in scores):
                    raise ValueError(f"Invalid regional signal: {path}")
                maxima.append(max(scores))
    return {
        "h3AcceptedFileCount": len(h3),
        "rnaAcceptedFileCount": len(rna),
        "sampleCount": len(samples),
        "expressionConditionRecordCount": len(conditions),
        "expressionReplicateRecordCount": len(replicates),
        "regionalBigWigFileCount": len(h3),
        "regionalBigWigTotalBytes": total_size,
        "maximumRegionalSignal": max(maxima),
    }


def artifact_identities(h3: list[FileLock]) -> dict[str, dict[str, Any]]:
    paths = [BIGWIG_DIR / f"{lock.sample_id}.bigWig" for lock in h3]
    paths.extend(
        OUTPUT_DIR / name
        for name in (
            "samples.tsv",
            "expression-conditions.tsv",
            "expression-replicates.tsv",
            "regions.tsv",
            "elements.tsv",
            "genes.tsv",
            "candidate-assessment.tsv",
            "selection-report.tsv",
        )
    )
    return {
        path.relative_to(RECIPE_DIR).as_posix(): {
            "fileSizeBytes": path.stat().st_size,
            "sha256": digest(path),
        }
        for path in sorted(paths)
    }


def main() -> None:
    args = parse_args()
    provenance = load_provenance()
    h3 = file_locks(source(provenance, "encode-h3k27ac"))
    rna = file_locks(source(provenance, "encode-rna"))
    retained_regions = regions(provenance)
    selected_genes = genes(provenance)
    if len(h3) != 24 or len(rna) != 24:
        raise ValueError("The accepted panel must contain 24 ChIP and 24 RNA files")
    if not args.verify_only:
        gtf_path, supplement_path = prepare_inputs(provenance, rna)
        validate_supplement(supplement_path)
        annotation_rows = read_gene_annotation(
            gtf_path, retained_regions, selected_genes
        )
        rna_values = read_expression(rna, selected_genes)
        write_tables(
            provenance,
            h3,
            rna_values,
            selected_genes,
            retained_regions,
            annotation_rows,
        )
        for index, lock in enumerate(h3, start=1):
            extract_bigwig(lock, retained_regions)
            print(f"Prepared regional BigWig {index}/24: {lock.accession}")
    validation = validate_outputs(h3, rna, selected_genes, retained_regions)
    artifacts = artifact_identities(h3)
    accepted = provenance.get("distribution", {}).get("artifacts")
    if accepted is not None and accepted != artifacts:
        raise ValueError("Generated outputs do not match committed artifact identities")
    accepted_run = {
        "distribution": {
            "baseUrl": (
                "https://data.genomespy.app/datasets/"
                f"encode-mouse-fetal-development-mm10/{provenance['releaseId']}/"
            ),
            "artifacts": artifacts,
        },
        "outputs": {
            "regionalBigWigs": {
                "path": "output/bigwigs/{sampleId}.bigWig",
                "fileCount": 24,
                "totalFileSizeBytes": validation["regionalBigWigTotalBytes"],
            },
            "samples": {"path": "output/samples.tsv", "recordCount": 24},
            "expressionConditions": {
                "path": "output/expression-conditions.tsv",
                "recordCount": 96,
            },
            "expressionReplicates": {
                "path": "output/expression-replicates.tsv",
                "recordCount": 192,
            },
            "regions": {"path": "output/regions.tsv", "recordCount": 4},
            "elements": {"path": "output/elements.tsv", "recordCount": len(ELEMENTS)},
            "genes": {
                "path": "output/genes.tsv",
                "recordCount": len(read_tsv(OUTPUT_DIR / "genes.tsv")),
            },
            "candidateAssessment": {
                "path": "output/candidate-assessment.tsv",
                "recordCount": len(read_tsv(OUTPUT_DIR / "candidate-assessment.tsv")),
            },
            "selectionReport": {
                "path": "output/selection-report.tsv",
                "recordCount": len(read_tsv(OUTPUT_DIR / "selection-report.tsv")),
            },
        },
        "validation": {
            **validation,
            "oneSignalFilePerBiologicalReplicate": True,
            "pooledAndPseudoreplicateSignalsExcluded": True,
            "numericStageOrdering": True,
            "expressionMeansTraceToTwoRnaReplicates": True,
            "expressionZScoresStandardizedAcrossTissueStagePanel": True,
            "sampleExpressionValuesMatchConditionMeans": True,
            "regionalBigWigStructureAndBoundsValid": True,
            "allObservedSignalsFiniteAndNonnegative": True,
            "supplementaryRecordsMatchPinnedWorkbook": True,
            "sourceBigWigIdentityCheck": (
                "ENCODE-reported sizes and MD5 values are pinned. Remote headers "
                "and extracted intervals are read from pinned public object URLs; "
                "complete upstream BigWig payloads are not downloaded or "
                "checksummed locally."
            ),
        },
    }
    write_text(
        ACCEPTED_RUN_PATH, json.dumps(accepted_run, indent=2, sort_keys=True) + "\n"
    )
    action = "Verified" if args.verify_only else "Prepared and verified"
    print(
        f"{action} 24 H3K27ac rows, 24 RNA inputs, and four regional extracts per row."
    )
    print(f"Accepted-run metadata: {ACCEPTED_RUN_PATH}")


if __name__ == "__main__":
    main()
