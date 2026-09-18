# mypy: ignore-errors
"""Wrangle MCCA workbooks into GenomeSpy-ready tabular data.

This module handles the source XLSX files for cell line annotations, copy-ratio
segments, somatic mutation annotations, and model alleles parsed from
`MouseModelDetailed`. It preserves the original MCCA metadata column names,
filters genomic rows to canonical mm10 chromosomes, normalizes common missing
value tokens, and writes typed Parquet files for GenomeSpy. Transcriptome
expression has a separate Zarr-specific path in `expression.py`.
"""

from __future__ import annotations

import argparse
from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from mcca_genomespy.expression import (
    load_gencode_gene_map,
    write_transcriptome_zarr,
)
from mcca_genomespy.sequencing_metadata import (
    LCWGS_RUN_REPORT,
    WES_RUN_REPORT,
    read_run_report_sample_ids,
    sequencing_metadata_rows,
    write_sequencing_metadata_parquet,
)
from mcca_genomespy.values import is_missing_value

CANONICAL_CHROMS = tuple([f"chr{chrom}" for chrom in range(1, 20)] + ["chrX", "chrY"])
CANONICAL_CHROM_SET = set(CANONICAL_CHROMS)

METADATA_SCHEMA = {
    "MCCA-ID": pa.string(),
    "CellLineName": pa.string(),
    "MouseID": pa.string(),
    "TumorLocation": pa.string(),
    "CellLineSource": pa.string(),
    "CellLineDistributor": pa.string(),
    "PublicationStatus": pa.string(),
    "MouseModelType": pa.string(),
    "MouseModel": pa.string(),
    "MouseModelDetailed": pa.string(),
    "Tissue": pa.string(),
    "Lineage": pa.string(),
    "Site": pa.string(),
    "CancerType": pa.string(),
    "CancerTypeDetailed": pa.string(),
    "MicroscopicMorphology": pa.string(),
    "MicroscopicMorphologyDetailed": pa.string(),
    "SurvivalDays": pa.int32(),
    "DistantMetastasis": pa.string(),
    "ComplexRearrangement": pa.string(),
    "Chromothripsis": pa.string(),
    "Gender": pa.string(),
    "ImmunocompetentTransplantation": pa.string(),
}

CNV_SCHEMA = {
    "sample": pa.string(),
    "chrom": pa.string(),
    "start": pa.int32(),
    "end": pa.int32(),
    "log2fc": pa.float64(),
}
MUTATION_SCHEMA = {
    "sample": pa.string(),
    "chrom": pa.string(),
    "pos": pa.int32(),
    "ref": pa.string(),
    "alt": pa.string(),
    "tumor_af": pa.float64(),
    "tumor_ref_depth": pa.int32(),
    "tumor_alt_depth": pa.int32(),
    "normal_ref_depth": pa.int32(),
    "normal_alt_depth": pa.int32(),
    "gene": pa.string(),
    "effect": pa.string(),
    "impact": pa.string(),
    "feature_id": pa.string(),
    "hgvs_c": pa.string(),
    "hgvs_p": pa.string(),
    "normal_ngs": pa.string(),
}
TRANSCRIPTOME_ZIP = "MCCA-Transcriptomes-VsdBatchCorrected-2025Q2.zip"
GENCODE_GTF = "gencode.vM25.annotation.gtf.gz"
MODEL_ALLELE_SAMPLE_ID_FIELD = "MCCA-ID"
MODEL_ALLELE_EXCLUDED_SYMBOLS = {
    "Bacterial",
    "Chemical",
    "Hormone",
    "Irradiation",
    "Resistance",
    "Viral",
}


def normalize_chrom(value: Any) -> str | None:
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text.startswith("chr"):
        text = "chr" + text

    if text in CANONICAL_CHROM_SET:
        return text
    return None


def parse_int(value: Any) -> int:
    return int(float(value))


def parse_float(value: Any) -> float:
    return float(value)


def parse_optional_int(value: Any) -> int | None:
    if is_missing_value(value):
        return None
    return parse_int(value)


def clean_text(value: Any) -> str:
    if is_missing_value(value):
        return ""
    return str(value).strip()


def clean_metadata_value(value: Any) -> str | int | None:
    if is_missing_value(value):
        return None
    return str(value).strip()


def clean_metadata_numeric_value(value: Any) -> int | None:
    if is_missing_value(value):
        return None
    return parse_int(value)


def parse_model_alleles(value: Any) -> dict[str, str]:
    if is_missing_value(value):
        return {}

    alleles: dict[str, list[str]] = {}
    for component in str(value).split(";"):
        component = component.strip()
        if "," not in component:
            continue

        symbol, allele = (part.strip() for part in component.split(",", 1))
        if (
            symbol in MODEL_ALLELE_EXCLUDED_SYMBOLS
            or symbol.startswith("rAAV")
            or is_missing_value(symbol)
            or is_missing_value(allele)
        ):
            continue

        alleles.setdefault(symbol, []).append(allele)

    return {symbol: ";".join(values) for symbol, values in alleles.items()}


def read_xlsx_rows(path: Path) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    header = [clean_text(value) for value in next(rows)]

    for row in rows:
        if any(value is not None for value in row):
            yield dict(zip(header, row, strict=False))


def normalize_metadata_row(row: dict[str, Any]) -> dict[str, str | int | None]:
    return {
        column: clean_metadata_numeric_value(row[column])
        if pa.types.is_integer(schema_type)
        else clean_metadata_value(row[column])
        for column, schema_type in METADATA_SCHEMA.items()
    }


def normalize_cnv_row(row: dict[str, Any]) -> dict[str, Any] | None:
    chrom = normalize_chrom(row["CHROM"])
    if chrom is None:
        return None

    return {
        "sample": clean_text(row["MCCA_ID"]),
        "chrom": chrom,
        "start": parse_int(row["START-mm10"]),
        "end": parse_int(row["END-mm10"]),
        "log2fc": parse_float(row["LOG2FC"]),
    }


def normalize_mutation_row(row: dict[str, Any]) -> dict[str, Any] | None:
    chrom = normalize_chrom(row["CHROM"])
    if chrom is None:
        return None

    return {
        "sample": clean_text(row["MCCA_ID"]),
        "chrom": chrom,
        "pos": parse_int(row["POS-mm10"]),
        "ref": clean_text(row["REF"]),
        "alt": clean_text(row["ALT"]),
        "tumor_af": parse_float(row["GEN[Tumor].AF"]),
        "tumor_ref_depth": parse_optional_int(row["GEN[Tumor].AD[0]"]),
        "tumor_alt_depth": parse_optional_int(row["GEN[Tumor].AD[1]"]),
        "normal_ref_depth": parse_optional_int(row["GEN[Normal].AD[0]"]),
        "normal_alt_depth": parse_optional_int(row["GEN[Normal].AD[1]"]),
        "gene": clean_text(row["ANN[*].GENE"]),
        "effect": clean_text(row["ANN[*].EFFECT"]),
        "impact": clean_text(row["ANN[*].IMPACT"]),
        "feature_id": clean_text(row["ANN[*].FEATUREID"]),
        "hgvs_c": clean_text(row["ANN[*].HGVS_C"]),
        "hgvs_p": clean_text(row["ANN[*].HGVS_P"]),
        "normal_ngs": clean_text(row["NORMAL_NGS"]),
    }


def write_parquet(
    path: Path,
    schema: dict[str, pa.DataType],
    rows: Iterable[dict[str, Any]],
) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(schema)
    columns = {field: [] for field in fieldnames}
    count = 0

    for row in rows:
        for field in fieldnames:
            columns[field].append(row[field])
        count += 1

    arrays = [pa.array(columns[field], type=schema[field]) for field in fieldnames]
    table = pa.Table.from_arrays(arrays, names=fieldnames)
    pq.write_table(table, path, compression="snappy")

    return count


def write_metadata_parquet(path: Path, rows: Iterable[dict[str, Any]]) -> int:
    return write_parquet(path, METADATA_SCHEMA, rows)


def model_allele_rows(rows: Iterable[dict[str, Any]]) -> Iterable[dict[str, str]]:
    for row in rows:
        alleles = parse_model_alleles(row["MouseModelDetailed"])
        yield {
            MODEL_ALLELE_SAMPLE_ID_FIELD: row["MCCA-ID"],
            **alleles,
        }


def write_model_alleles_parquet(path: Path, rows: Iterable[dict[str, str]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    row_list = list(rows)
    symbols = sorted(
        {
            key
            for row in row_list
            for key in row.keys()
            if key != MODEL_ALLELE_SAMPLE_ID_FIELD
        }
    )
    columns = [MODEL_ALLELE_SAMPLE_ID_FIELD, *symbols]

    arrays = [
        pa.array(
            [row.get(column) for row in row_list],
            type=pa.string(),
        )
        for column in columns
    ]
    table = pa.Table.from_arrays(arrays, names=columns)
    pq.write_table(table, path, compression="snappy")
    return len(row_list)


def normalized_rows(
    rows: Iterable[dict[str, Any]],
    normalizer: Any,
) -> Iterable[dict[str, Any]]:
    for row in rows:
        normalized = normalizer(row)
        if normalized is not None:
            yield normalized


def wrangle(raw_dir: Path, processed_dir: Path) -> None:
    raw_metadata_rows = list(read_xlsx_rows(raw_dir / "cell_line_annotations.xlsx"))
    metadata_rows = list(normalized_rows(raw_metadata_rows, normalize_metadata_row))

    metadata_count = write_metadata_parquet(
        processed_dir / "samples.parquet",
        metadata_rows,
    )
    model_allele_count = write_model_alleles_parquet(
        processed_dir / "model-alleles.parquet",
        model_allele_rows(metadata_rows),
    )
    sequencing_metadata_count = write_sequencing_metadata_parquet(
        processed_dir / "sequencing.parquet",
        sequencing_metadata_rows(
            metadata_rows,
            read_run_report_sample_ids(raw_dir / LCWGS_RUN_REPORT),
            read_run_report_sample_ids(raw_dir / WES_RUN_REPORT),
        ),
    )
    cnv_count = write_parquet(
        processed_dir / "copy-ratios.parquet",
        CNV_SCHEMA,
        normalized_rows(
            read_xlsx_rows(raw_dir / "copy_number_variation.xlsx"),
            normalize_cnv_row,
        ),
    )
    mutation_count = write_parquet(
        processed_dir / "mutations.parquet",
        MUTATION_SCHEMA,
        normalized_rows(
            read_xlsx_rows(raw_dir / "mutations.xlsx"),
            normalize_mutation_row,
        ),
    )

    print(f"Wrote {metadata_count} samples")
    print(f"Wrote {model_allele_count} model allele metadata rows")
    print(f"Wrote {sequencing_metadata_count} sequencing metadata rows")
    print(f"Wrote {cnv_count} canonical copy-ratio segments")
    print(f"Wrote {mutation_count} canonical mutation annotations")

    transcriptome_path = raw_dir / TRANSCRIPTOME_ZIP
    if transcriptome_path.exists():
        gene_map = load_gencode_gene_map(raw_dir / GENCODE_GTF)
        summary = write_transcriptome_zarr(
            transcriptome_path,
            processed_dir / "expression.zarr",
            gene_map=gene_map,
        )
        print(
            "Wrote transcriptome Zarr with "
            + str(summary["samples"])
            + " samples and "
            + str(summary["genes"])
            + " genes"
        )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Wrangle MCCA source workbooks into GenomeSpy-ready data files."
    )
    parser.add_argument(
        "--raw-dir",
        type=Path,
        default=Path(__file__).resolve().parents[2] / "download",
        help="Directory containing downloaded source workbooks.",
    )
    parser.add_argument(
        "--processed-dir",
        type=Path,
        default=Path(__file__).resolve().parents[2] / "output" / "processed",
        help="Directory for generated GenomeSpy-ready data files.",
    )
    args = parser.parse_args()

    wrangle(args.raw_dir, args.processed_dir)


if __name__ == "__main__":
    main()
